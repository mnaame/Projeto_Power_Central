"""Telas do ajuste em lote — o que protege é quem PODE apertar o botão."""

from io import BytesIO

from openpyxl import load_workbook

from app.extensions import db
from app.services import horario_lote_service as svc
from app.services import settings_service
from tests.integration.test_horario_lote_service import FakePortal


def _lote(app, portal=None):
    lote = svc.preparar(config=app.config, softguard_client=portal or FakePortal())
    db.session.commit()
    return lote


# ---------- quem pode ----------


def test_index_exige_login(client):
    assert client.get("/horario-lote").status_code == 302


def test_operador_nao_entra(operador_client):
    """Escreve no cadastro de todas as contas — não é tela de operador."""
    assert operador_client.get("/horario-lote").status_code == 403


def test_operador_nao_prepara(operador_client):
    assert operador_client.post("/horario-lote/preparar").status_code == 403


def test_operador_nao_aplica(app, operador_client):
    lote = _lote(app)
    assert operador_client.post(f"/horario-lote/{lote.id}/aplicar").status_code == 403


def test_admin_entra(admin_client):
    assert admin_client.get("/horario-lote").status_code == 200


# ---------- confirmação ----------


def test_lote_inteiro_sem_a_palavra_de_confirmacao_nao_executa(app, admin_client, monkeypatch):
    settings_service.set("horario_lote_simulacao", "false")
    db.session.commit()
    portal = FakePortal()
    lote = _lote(app, portal)
    monkeypatch.setattr(svc, "_criar_cliente", lambda config: portal)

    resposta = admin_client.post(
        f"/horario-lote/{lote.id}/aplicar", data={}, follow_redirects=True
    )

    assert portal.puts == []
    assert b"APLICAR" in resposta.data


def test_lote_inteiro_com_a_confirmacao_executa(app, admin_client, monkeypatch):
    settings_service.set("horario_lote_simulacao", "false")
    db.session.commit()
    portal = FakePortal()
    lote = _lote(app, portal)
    monkeypatch.setattr(svc, "_criar_cliente", lambda config: portal)

    admin_client.post(
        f"/horario-lote/{lote.id}/aplicar",
        data={"confirmacao": "APLICAR"}, follow_redirects=True,
    )

    assert len(portal.puts) == 2


def test_teste_em_uma_conta_dispensa_a_palavra(app, admin_client, monkeypatch):
    """O passo barato e seguro tem que ser fácil de dar — senão ninguém dá
    e vai direto para o lote."""
    settings_service.set("horario_lote_simulacao", "false")
    db.session.commit()
    portal = FakePortal()
    lote = _lote(app, portal)
    monkeypatch.setattr(svc, "_criar_cliente", lambda config: portal)

    admin_client.post(
        f"/horario-lote/{lote.id}/aplicar",
        data={"apenas_conta": "0118"}, follow_redirects=True,
    )

    assert [str(p["Id"]) for p in portal.puts] == ["7201"]


# ---------- backup e listas ----------


def test_backup_tem_o_estado_antes(app, admin_client):
    lote = _lote(app)

    resposta = admin_client.get(f"/horario-lote/{lote.id}/backup")

    assert resposta.status_code == 200
    aba = load_workbook(BytesIO(resposta.data))["ANTES"]
    assert [c.value for c in aba[1]][:2] == ["CONTA", "NOME"]
    fechamentos = {aba.cell(row=linha, column=7).value for linha in range(2, 5)}
    assert fechamentos == {"18:00", "23:00", "19:30"}


def test_exportar_separa_aplicar_de_manual(app, admin_client):
    lote = _lote(app)

    resposta = admin_client.get(f"/horario-lote/{lote.id}/exportar")

    wb = load_workbook(BytesIO(resposta.data))
    assert wb.sheetnames == ["APLICAR", "MANUAL (vira 00h)"]
    assert wb["APLICAR"].max_row == 3  # cabeçalho + 2
    assert wb["MANUAL (vira 00h)"].max_row == 2  # cabeçalho + 1


def test_tela_mostra_as_duas_listas(app, admin_client):
    _lote(app)

    corpo = admin_client.get("/horario-lote").get_data(as_text=True)

    assert "18:00 → 19:00" in corpo
    assert "Manual — vira 00h" in corpo
    assert "SHOPPING VETTORE" in corpo


def test_tela_avisa_quando_a_simulacao_esta_desligada(app, admin_client):
    settings_service.set("horario_lote_simulacao", "false")
    db.session.commit()

    corpo = admin_client.get("/horario-lote").get_data(as_text=True)

    assert "SIMULAÇÃO DESLIGADA" in corpo


def test_marcar_alterar_de_verdade_grava_mesmo_em_simulacao(app, admin_client, monkeypatch):
    """O caminho que estava bloqueado: testar numa conta sem desarmar a
    chave global, que armaria o botão do lote inteiro junto."""
    portal = FakePortal()
    lote = _lote(app, portal)
    monkeypatch.setattr(svc, "_criar_cliente", lambda config: portal)

    corpo = admin_client.post(
        f"/horario-lote/{lote.id}/aplicar",
        data={"apenas_conta": "118", "gravar_de_verdade": "1"},
        follow_redirects=True,
    ).get_data(as_text=True)

    assert [str(p["Id"]) for p in portal.puts] == ["7201"]
    assert "Alterado na conta 118" in corpo


def test_alterar_de_verdade_sem_conta_nao_grava_o_lote(app, admin_client, monkeypatch):
    """A caixa libera UMA conta. Sem conta preenchida ela não vale, e o
    lote inteiro ainda exige a palavra de confirmação."""
    portal = FakePortal()
    lote = _lote(app, portal)
    monkeypatch.setattr(svc, "_criar_cliente", lambda config: portal)

    admin_client.post(
        f"/horario-lote/{lote.id}/aplicar",
        data={"gravar_de_verdade": "1", "confirmacao": "APLICAR"},
        follow_redirects=True,
    )

    assert portal.puts == []


def test_conta_inexistente_avisa_na_tela(app, admin_client, monkeypatch):
    portal = FakePortal()
    lote = _lote(app, portal)
    monkeypatch.setattr(svc, "_criar_cliente", lambda config: portal)

    corpo = admin_client.post(
        f"/horario-lote/{lote.id}/aplicar",
        data={"apenas_conta": "9999"}, follow_redirects=True,
    ).get_data(as_text=True)

    assert "9999" in corpo
    assert "Manual" in corpo  # diz onde a conta pode estar
