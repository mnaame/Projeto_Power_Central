from io import BytesIO

from openpyxl import load_workbook

from app.domain.horarios import ContaAuditada
from app.extensions import db
from app.services import auditoria_horarios_service
from tests.integration.test_auditoria_horarios_service import (
    FAIXA,
    FakeSoftGuardClient,
    _conta,
)


def _semear_snapshot():
    """Grava um snapshot sem tocar no portal — a tela e o export leem
    daqui, não da varredura."""
    resultado = {
        "sem": [ContaAuditada(conta="96", nome="LOJA SEM HORARIO")],
        "com": [
            ContaAuditada(
                conta="95", nome="LOJA COM HORARIO", resumo="1 faixa(s) 06:00–18:00"
            )
        ],
        "erros": 0,
        "total": 2,
    }
    auditoria_horarios_service.salvar_snapshot(resultado)
    db.session.commit()


def test_index_exige_login(client):
    resposta = client.get("/auditoria-horarios")
    assert resposta.status_code == 302
    assert "/login" in resposta.headers["Location"]


def test_rodar_exige_login(client):
    resposta = client.post("/auditoria-horarios/rodar")
    assert resposta.status_code == 302
    assert "/login" in resposta.headers["Location"]


def test_exportar_exige_login(client):
    resposta = client.get("/auditoria-horarios/exportar")
    assert resposta.status_code == 302
    assert "/login" in resposta.headers["Location"]


def test_index_sem_auditoria_mostra_convite(operador_client):
    resposta = operador_client.get("/auditoria-horarios")
    assert resposta.status_code == 200
    assert "Nenhuma auditoria rodada ainda" in resposta.get_data(as_text=True)


def test_index_mostra_as_duas_listas_do_snapshot(operador_client, app):
    _semear_snapshot()

    corpo = operador_client.get("/auditoria-horarios").get_data(as_text=True)

    assert "LOJA SEM HORARIO" in corpo
    assert "LOJA COM HORARIO" in corpo
    assert "06:00–18:00" in corpo


def test_exportar_sem_auditoria_da_404(operador_client):
    assert operador_client.get("/auditoria-horarios/exportar").status_code == 404


def test_export_gera_as_duas_abas(operador_client, app):
    _semear_snapshot()

    resposta = operador_client.get("/auditoria-horarios/exportar")
    assert resposta.status_code == 200

    wb = load_workbook(BytesIO(resposta.data))
    assert wb.sheetnames == ["SEM horário", "COM horário"]

    aba_sem = wb["SEM horário"]
    assert [c.value for c in aba_sem[1]] == ["CONTA", "NOME"]
    assert [c.value for c in aba_sem[2]] == ["96", "LOJA SEM HORARIO"]

    aba_com = wb["COM horário"]
    assert [c.value for c in aba_com[1]] == ["CONTA", "NOME", "RESUMO"]
    assert aba_com[2][2].value == "1 faixa(s) 06:00–18:00"


def test_rodar_grava_snapshot_e_redireciona(operador_client, app, monkeypatch):
    """POST/Redirect/GET: o resultado fica salvo e um F5 não dispara a
    varredura de novo."""
    fake = FakeSoftGuardClient(
        contas=[_conta("10", "95", "LOJA A"), _conta("11", "96", "LOJA B")],
        horarios={"10": [], "11": FAIXA},
    )
    monkeypatch.setattr(
        auditoria_horarios_service, "_criar_cliente", lambda config: fake
    )

    resposta = operador_client.post("/auditoria-horarios/rodar")
    assert resposta.status_code == 302
    assert resposta.headers["Location"].endswith("/auditoria-horarios")

    snapshot = auditoria_horarios_service.ultimo_snapshot()
    assert snapshot.sem == 1
    assert snapshot.com == 1


def test_rodar_com_portal_fora_avisa_sem_derrubar_a_tela(operador_client, app, monkeypatch):
    from app.integrations.softguard_client import SoftGuardError

    fake = FakeSoftGuardClient(contas=SoftGuardError("portal fora do ar"))
    monkeypatch.setattr(
        auditoria_horarios_service, "_criar_cliente", lambda config: fake
    )

    resposta = operador_client.post("/auditoria-horarios/rodar", follow_redirects=True)

    assert resposta.status_code == 200
    assert "portal fora do ar" in resposta.get_data(as_text=True)
    assert auditoria_horarios_service.ultimo_snapshot() is None


def test_dashboard_mostra_o_card_sem_rodar_a_varredura(operador_client, app, monkeypatch):
    """O card lê do snapshot. Se ele chamasse `auditar`, este teste
    explodiria — é essa a garantia que interessa."""
    def _explode(*args, **kwargs):  # pragma: no cover
        raise AssertionError("o dashboard não pode rodar a varredura")

    monkeypatch.setattr(auditoria_horarios_service, "auditar", _explode)
    _semear_snapshot()

    corpo = operador_client.get("/").get_data(as_text=True)

    assert "Saúde do cadastro" in corpo
    assert "1 conta(s) sem horário" in corpo


def test_dashboard_sem_snapshot_convida_a_rodar(operador_client, app):
    corpo = operador_client.get("/").get_data(as_text=True)
    assert "Rode a auditoria para saber" in corpo
