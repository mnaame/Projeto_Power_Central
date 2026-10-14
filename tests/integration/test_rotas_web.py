from io import BytesIO

from openpyxl import load_workbook

from app.services import rotas_service as svc
from tests.integration.test_rotas_service import (
    ENDERECO_BURITIS,
    ENDERECO_SAVASSI,
    FakeAuvo,
    FakeSoftGuard,
    _cliente_auvo,
    _conta,
    _xlsx_auvo,
)

CONTAS = [
    _conta("1", "A"),
    _conta("2", "B"),
    _conta("3", "C"),
    _conta("4", "SEM BAIRRO", cidade="CONTAGEM"),
]
CLIENTES = [
    _cliente_auvo(1, "A", ENDERECO_SAVASSI),
    _cliente_auvo(2, "B", ENDERECO_SAVASSI),
    _cliente_auvo(3, "C", ENDERECO_BURITIS),
]


def _preparar(monkeypatch, contas=None, clientes=None):
    monkeypatch.setattr(
        svc, "SoftGuardClient", lambda *a, **k: FakeSoftGuard(contas or CONTAS)
    )
    monkeypatch.setattr(svc, "credenciais_softguard", lambda config: None)
    monkeypatch.setattr(
        svc.auvo_service,
        "criar_cliente",
        lambda config: FakeAuvo(CLIENTES if clientes is None else clientes),
    )


def test_index_exige_login(client):
    resposta = client.get("/rotas")
    assert resposta.status_code == 302
    assert "/login" in resposta.headers["Location"]


def test_operador_pode_ver(operador_client):
    """Só leitura nas plataformas — não é tela de admin."""
    assert operador_client.get("/rotas").status_code == 200


def test_index_sem_catalogo_convida_a_gerar(operador_client):
    corpo = operador_client.get("/rotas").get_data(as_text=True)

    assert "Nenhum catálogo gerado ainda" in corpo


def test_gerar_mostra_as_regioes_e_os_clientes(operador_client, monkeypatch):
    _preparar(monkeypatch)

    corpo = operador_client.post("/rotas/gerar", follow_redirects=True).get_data(
        as_text=True
    )

    assert "Savassi — Belo Horizonte" in corpo
    assert "Buritis — Belo Horizonte" in corpo
    assert "SEM BAIRRO" in corpo  # quem não tem bairro não some da tela


def test_catalogo_sobrevive_a_um_navegador_novo(operador_client, monkeypatch):
    """O catálogo vivia no cookie de sessão. Com a base real ele dava
    4.746 bytes contra um limite de 4.093 no navegador: o cookie era
    descartado calado e a tela respondia "nenhum catálogo gerado ainda"
    logo depois de gerar, sem erro nenhum. Agora mora no banco, então um
    cliente HTTP sem nenhum cookie tem que ver o mesmo catálogo."""
    _preparar(monkeypatch)
    operador_client.post("/rotas/gerar")

    operador_client.delete_cookie("session")
    corpo = operador_client.get("/rotas").get_data(as_text=True)

    assert "Savassi — Belo Horizonte" in corpo


def test_catalogo_grande_nao_depende_do_tamanho_do_cookie(operador_client, monkeypatch):
    """O caso que quebrou em produção: base inteira, não quatro contas."""
    contas = [_conta(str(i), f"CLIENTE NUMERO {i}") for i in range(1, 120)]
    clientes = [
        _cliente_auvo(i, f"CLIENTE NUMERO {i}",
                      f"Rua {i}, {i},,Bairro {i},Belo Horizonte,MG,3010000{i % 10}")
        for i in range(1, 120)
    ]
    _preparar(monkeypatch, contas=contas, clientes=clientes)

    operador_client.post("/rotas/gerar")
    corpo = operador_client.get("/rotas").get_data(as_text=True)

    assert "CLIENTE NUMERO 119" in corpo


def test_exportar_sem_gerar_avisa(operador_client):
    corpo = operador_client.get("/rotas/exportar", follow_redirects=True).get_data(
        as_text=True
    )

    assert "Gere o catálogo antes" in corpo


def test_export_tem_as_tres_abas(operador_client, monkeypatch):
    _preparar(monkeypatch)
    operador_client.post("/rotas/gerar")

    resposta = operador_client.get("/rotas/exportar")

    assert resposta.status_code == 200
    wb = load_workbook(BytesIO(resposta.data))
    assert wb.sheetnames == ["Clientes por bairro", "Resumo por região", "Sem bairro"]
    assert [c.value for c in wb["Clientes por bairro"][1]][:4] == [
        "REGIÃO", "CONTA", "CLIENTE", "BAIRRO",
    ]
    assert wb["Sem bairro"].max_row == 2  # cabeçalho + o cliente sem bairro


# ---------- marcar visita ----------


def _marcar(client, conta, visitado="1"):
    return client.post(
        "/rotas/visita",
        data={"conta": conta, "visitado": visitado},
        headers={"Accept": "application/json"},
    )


def test_marcar_visita_registra_e_devolve_o_progresso(operador_client, monkeypatch):
    _preparar(monkeypatch)
    operador_client.post("/rotas/gerar")

    dados = _marcar(operador_client, "1").get_json()

    assert dados["visitado"] is True
    assert dados["visitados"] == 1
    assert dados["total"] == 4


def test_desmarcar_desfaz(operador_client, monkeypatch):
    """Marcar errado é fácil e inevitável; sem o desfazer, o único
    conserto seria reiniciar a rodada e perder todas as outras marcas."""
    _preparar(monkeypatch)
    operador_client.post("/rotas/gerar")
    _marcar(operador_client, "1")

    dados = _marcar(operador_client, "1", visitado="0").get_json()

    assert dados["visitado"] is False
    assert dados["visitados"] == 0


def test_visita_aparece_na_tela_depois_de_recarregar(operador_client, monkeypatch):
    _preparar(monkeypatch)
    operador_client.post("/rotas/gerar")
    _marcar(operador_client, "1")

    corpo = operador_client.get("/rotas").get_data(as_text=True)

    assert "1 de 4 cliente(s) visitado(s)" in corpo
    assert "is-visitado" in corpo


def test_marcar_sem_javascript_tambem_funciona(operador_client, monkeypatch):
    """O `form` é de verdade: sem JS a tela recarrega, mas marca."""
    _preparar(monkeypatch)
    operador_client.post("/rotas/gerar")

    corpo = operador_client.post(
        "/rotas/visita", data={"conta": "1", "visitado": "1"}, follow_redirects=True
    ).get_data(as_text=True)

    assert "1 de 4 cliente(s) visitado(s)" in corpo


def test_marcar_conta_que_nao_esta_no_ciclo_nao_passa(operador_client, monkeypatch):
    _preparar(monkeypatch)
    operador_client.post("/rotas/gerar")

    assert _marcar(operador_client, "9999").status_code == 404


def test_gerar_de_novo_preserva_as_visitas(operador_client, monkeypatch):
    """Gerar no meio da rodada é normal (planilha nova, conta nova), e
    perder as marcas faria o técnico repetir locais — justamente o que a
    marcação existe para evitar."""
    _preparar(monkeypatch)
    operador_client.post("/rotas/gerar")
    _marcar(operador_client, "1")

    operador_client.post("/rotas/gerar")
    corpo = operador_client.get("/rotas").get_data(as_text=True)

    assert "1 de 4 cliente(s) visitado(s)" in corpo


# ---------- reiniciar o ciclo ----------


def test_reiniciar_zera_as_marcas_e_abre_o_proximo_ciclo(operador_client, monkeypatch):
    _preparar(monkeypatch)
    operador_client.post("/rotas/gerar")
    _marcar(operador_client, "1")
    _marcar(operador_client, "2")

    corpo = operador_client.post("/rotas/reiniciar", follow_redirects=True).get_data(
        as_text=True
    )

    assert "Ciclo #2" in corpo
    assert "0 de 4 cliente(s) visitado(s)" in corpo


def test_reiniciar_mantem_os_mesmos_clientes_sem_consultar_a_plataforma(
    app, operador_client, monkeypatch
):
    """Refazer a varredura só para zerar marcas custaria minutos para não
    mudar nada."""
    _preparar(monkeypatch)
    operador_client.post("/rotas/gerar")

    # A plataforma "some" depois de gerar: reiniciar não pode precisar dela.
    monkeypatch.setattr(
        svc, "SoftGuardClient", lambda *a, **k: (_ for _ in ()).throw(AssertionError())
    )
    corpo = operador_client.post("/rotas/reiniciar", follow_redirects=True).get_data(
        as_text=True
    )

    assert "SEM BAIRRO" in corpo


def test_aviso_de_rodada_concluida_so_aparece_com_tudo_marcado(
    operador_client, monkeypatch
):
    _preparar(monkeypatch)
    operador_client.post("/rotas/gerar")
    for conta in ("1", "2", "3"):
        _marcar(operador_client, conta)

    parcial = operador_client.get("/rotas").get_data(as_text=True)
    _marcar(operador_client, "4")
    completo = operador_client.get("/rotas").get_data(as_text=True)

    assert 'id="rotas-ciclo-completo"' in parcial and "display: none;" in parcial
    assert "Rodada concluída" in completo


def test_reiniciar_sem_ciclo_avisa(operador_client):
    corpo = operador_client.post("/rotas/reiniciar", follow_redirects=True).get_data(
        as_text=True
    )

    assert "Não há ciclo para reiniciar" in corpo


# ---------- importação da planilha ----------


def test_importar_planilha_da_auvo_pela_tela(operador_client):
    arquivo = _xlsx_auvo(
        [["Código", "Nome", "Endereço"], ["500", "PADARIA DO ZE", ENDERECO_SAVASSI]]
    )

    corpo = operador_client.post(
        "/rotas/importar",
        data={"arquivo": (BytesIO(arquivo), "clientes.xlsx")},
        content_type="multipart/form-data",
        follow_redirects=True,
    ).get_data(as_text=True)

    assert "Planilha da Auvo importada" in corpo
    assert "1 com bairro" in corpo


def test_importar_sem_arquivo_avisa(operador_client):
    corpo = operador_client.post(
        "/rotas/importar",
        data={},
        content_type="multipart/form-data",
        follow_redirects=True,
    ).get_data(as_text=True)

    assert "Escolha o arquivo" in corpo


def test_arquivo_que_nao_e_planilha_e_recusado_com_explicacao(operador_client):
    corpo = operador_client.post(
        "/rotas/importar",
        data={"arquivo": (BytesIO(b"nao sou planilha"), "qualquer.xls")},
        content_type="multipart/form-data",
        follow_redirects=True,
    ).get_data(as_text=True)

    assert "HTML" in corpo  # explica o que fazer, não só "deu erro"


def test_importar_nao_apaga_o_ciclo_nem_as_visitas(operador_client, monkeypatch):
    """Importar planilha é rotina (cadastro mudou), e descartar o ciclo
    jogaria fora as visitas já marcadas. O catálogo fica como está até
    alguém gerar de novo — e a mensagem diz isso."""
    _preparar(monkeypatch)
    operador_client.post("/rotas/gerar")
    _marcar(operador_client, "1")

    arquivo = _xlsx_auvo(
        [["Código", "Nome", "Endereço"], ["500", "PADARIA", ENDERECO_SAVASSI]]
    )
    corpo = operador_client.post(
        "/rotas/importar",
        data={"arquivo": (BytesIO(arquivo), "clientes.xlsx")},
        content_type="multipart/form-data",
        follow_redirects=True,
    ).get_data(as_text=True)

    assert "Gere o catálogo de novo para usar" in corpo
    assert "1 de 4 cliente(s) visitado(s)" in corpo


def test_tela_mostra_o_que_ja_foi_importado(operador_client):
    arquivo = _xlsx_auvo(
        [["Código", "Nome", "Endereço"], ["500", "PADARIA", ENDERECO_SAVASSI]]
    )
    operador_client.post(
        "/rotas/importar",
        data={"arquivo": (BytesIO(arquivo), "clientes.xlsx")},
        content_type="multipart/form-data",
    )

    corpo = operador_client.get("/rotas").get_data(as_text=True)

    assert "Endereços importados" in corpo
    assert "Auvo" in corpo


# ---------- mapa, busca e correção de bairro na tela ----------


COM_COORD = [
    _cliente_auvo(1, "A", ENDERECO_SAVASSI),
    _cliente_auvo(2, "B", ENDERECO_SAVASSI),
    _cliente_auvo(3, "C", ENDERECO_BURITIS),
]
for _c in COM_COORD:
    _c["latitude"], _c["longitude"] = "-19,9191", "-43,9386"


def test_mapa_aparece_quando_ha_coordenada(operador_client, monkeypatch):
    _preparar(monkeypatch, clientes=COM_COORD)
    operador_client.post("/rotas/gerar")

    corpo = operador_client.get("/rotas").get_data(as_text=True)

    assert 'id="mapa-rotas"' in corpo
    assert "vendor/leaflet/leaflet.js" in corpo  # vendorizado, não CDN


def test_sem_coordenada_nenhuma_o_mapa_nao_carrega_o_leaflet(
    operador_client, monkeypatch
):
    """Carregar 190 KB de biblioteca para desenhar zero pinos seria só
    peso."""
    _preparar(monkeypatch, clientes=[])
    operador_client.post("/rotas/gerar")

    corpo = operador_client.get("/rotas").get_data(as_text=True)

    assert "leaflet.js" not in corpo


def test_cliente_sem_coordenada_nao_vira_pino_mas_fica_na_lista(
    operador_client, monkeypatch
):
    _preparar(monkeypatch, clientes=COM_COORD)
    operador_client.post("/rotas/gerar")

    corpo = operador_client.get("/rotas").get_data(as_text=True)

    assert "SEM BAIRRO" in corpo  # a conta 4, sem Auvo e sem coordenada
    assert "3 de 4 cliente(s) com coordenada" in corpo


def test_busca_esta_na_tela(operador_client, monkeypatch):
    _preparar(monkeypatch)
    operador_client.post("/rotas/gerar")

    corpo = operador_client.get("/rotas").get_data(as_text=True)

    assert 'id="rotas-busca"' in corpo


def test_corrigir_bairro_pela_tela(operador_client, monkeypatch):
    _preparar(monkeypatch)
    operador_client.post("/rotas/gerar")

    resposta = operador_client.post(
        "/rotas/bairro",
        data={"conta": "4", "bairro": "Eldorado", "cidade": "Contagem"},
        headers={"Accept": "application/json"},
    )

    assert resposta.get_json()["bairro"] == "Eldorado"
    corpo = operador_client.get("/rotas").get_data(as_text=True)
    assert "Eldorado — Contagem" in corpo


def test_corrigir_bairro_sem_javascript_tambem_funciona(operador_client, monkeypatch):
    _preparar(monkeypatch)
    operador_client.post("/rotas/gerar")

    corpo = operador_client.post(
        "/rotas/bairro",
        data={"conta": "4", "bairro": "Eldorado", "cidade": "Contagem"},
        follow_redirects=True,
    ).get_data(as_text=True)

    assert "bairro definido como Eldorado" in corpo


def test_correcao_de_bairro_sobrevive_a_gerar_de_novo(operador_client, monkeypatch):
    """Uma leitura automática não tem autoridade para desfazer o que
    alguém olhou e decidiu."""
    _preparar(monkeypatch)
    operador_client.post("/rotas/gerar")
    operador_client.post(
        "/rotas/bairro", data={"conta": "4", "bairro": "Eldorado", "cidade": "Contagem"}
    )

    operador_client.post("/rotas/gerar")
    corpo = operador_client.get("/rotas").get_data(as_text=True)

    assert "Eldorado — Contagem" in corpo


def test_tela_sugere_os_bairros_que_ja_existem(operador_client, monkeypatch):
    _preparar(monkeypatch)
    operador_client.post("/rotas/gerar")

    corpo = operador_client.get("/rotas").get_data(as_text=True)

    assert 'id="rotas-bairros"' in corpo
    assert '<option value="Savassi">' in corpo


def test_mapa_diz_quantos_tem_coordenada_mesmo_quando_da_zero(
    operador_client, monkeypatch
):
    """Esconder a seção quando dá zero já custou três rodadas nesta tela:
    sem o número à vista, "não tem mapa" e "não tem coordenada" são
    indistinguíveis."""
    _preparar(monkeypatch, clientes=[])
    operador_client.post("/rotas/gerar")

    corpo = operador_client.get("/rotas").get_data(as_text=True)

    assert "0 de 4 cliente(s) com coordenada" in corpo
    assert "Nenhum cliente com coordenada" in corpo


def test_planilha_antiga_sem_coordenada_e_apontada_como_causa(
    operador_client, monkeypatch
):
    """Planilha importada e nenhuma coordenada nela = enviada antes de o
    sistema ler esse campo. É a causa mais provável de um mapa vazio."""
    operador_client.post(
        "/rotas/importar",
        data={
            "arquivo": (
                BytesIO(
                    _xlsx_auvo(
                        [["Código", "Nome", "Endereço"], ["500", "A", ENDERECO_SAVASSI]]
                    )
                ),
                "clientes.xlsx",
            )
        },
        content_type="multipart/form-data",
    )
    _preparar(monkeypatch, clientes=[])
    operador_client.post("/rotas/gerar")

    corpo = operador_client.get("/rotas").get_data(as_text=True)

    assert "estão sem coordenada" in corpo
    assert "Importe as duas de novo" in corpo


def test_tabela_de_importacao_mostra_a_coluna_de_coordenada(operador_client):
    operador_client.post(
        "/rotas/importar",
        data={
            "arquivo": (
                BytesIO(
                    _xlsx_auvo(
                        [
                            ["Código", "Nome", "Endereço", "Latitude", "Longitude"],
                            ["500", "A", ENDERECO_SAVASSI, "-19,9191", "-43,9386"],
                        ]
                    )
                ),
                "clientes.xlsx",
            )
        },
        content_type="multipart/form-data",
    )

    corpo = operador_client.get("/rotas").get_data(as_text=True)

    assert "Com coordenada" in corpo


def test_busca_usa_o_estilo_de_campo_da_plataforma(operador_client, monkeypatch):
    """Solto, o input herda o estilo cru do navegador — que em tema
    escuro sai preto, fora do padrão de todas as outras telas."""
    _preparar(monkeypatch)
    operador_client.post("/rotas/gerar")

    corpo = operador_client.get("/rotas").get_data(as_text=True)
    inicio = corpo.index('id="rotas-busca"')
    trecho = corpo[inicio - 400:inicio]

    assert 'class="field"' in trecho


def test_mapa_carrega_o_agrupador_de_pinos(operador_client, monkeypatch):
    """Cem pinos crus em cima de BH viram uma mancha."""
    _preparar(monkeypatch, clientes=COM_COORD)
    operador_client.post("/rotas/gerar")

    corpo = operador_client.get("/rotas").get_data(as_text=True)

    assert "leaflet.markercluster.js" in corpo
    assert "MarkerCluster.Default.css" in corpo
